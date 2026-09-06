"""Hauptfenster: Drei-Schritte-Wizard für nicht-technische Anwender.

  Schritt 1: Excel-Datei wählen (Dialog oder Drag & Drop) + Blatt wählen
  Schritt 2: Spalte mit den Materialnummern ANKLICKEN (Vorschau der Tabelle)
  Schritt 3: Prüfung läuft – Fortschritt, Ampel-Liste, Pause/Abbruch

Der Orchestrator läuft im Worker-Thread; seine Callbacks werden über
Qt-Signale in den GUI-Thread gehoben.
"""
from __future__ import annotations

import logging
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QObject, QSettings, Qt, Signal
from PySide6.QtGui import QColor, QIcon, QPixmap
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QFileDialog, QHBoxLayout, QHeaderView,
    QLabel, QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMessageBox,
    QProgressBar, QPushButton, QSplitter, QStackedWidget, QTabWidget,
    QTableWidget, QTableWidgetItem, QTextEdit, QVBoxLayout, QWidget,
)

from ..core.models import (
    JobStatus, MaterialResult, RunConfig, Severity, SEVERITY_LABEL,
)
from ..core.orchestrator import Callbacks, Orchestrator, Progress
from ..report.excel_writer import read_materials
from ..sap.adapter import SapAdapter

log = logging.getLogger(__name__)

PREVIEW_ROWS = 50

STATUS_COLOR = {
    JobStatus.OK: QColor(70, 160, 70),
    JobStatus.FINDINGS: QColor(230, 145, 0),
    JobStatus.FAILED: QColor(200, 30, 30),
    JobStatus.SKIPPED: QColor(150, 150, 150),
}


def klartext(exc: Exception) -> str:
    """Technische Ausnahme in eine Anweisung für Anwender übersetzen.

    Nicht-technische Anwender können mit „PermissionError [Errno 13]"
    nichts anfangen – wohl aber mit „Die Datei ist noch in Excel geöffnet".
    """
    text = str(exc)
    if isinstance(exc, PermissionError) or "Errno 13" in text:
        return ("Die Datei ist gesperrt – vermutlich noch in Excel geöffnet. "
                "Bitte schließen und erneut versuchen.")
    if isinstance(exc, FileNotFoundError):
        return ("Die Datei wurde nicht gefunden. Wurde sie verschoben oder "
                "umbenannt?")
    if "No space left" in text or "Errno 28" in text:
        return ("Auf dem Laufwerk ist kein Platz mehr. Bitte Speicherplatz "
                "freigeben; der Lauf lässt sich danach fortsetzen.")
    if "pywin32" in text or "win32com" in text:
        return ("Die SAP-Anbindung fehlt (pywin32). Bitte an die "
                "Systembetreuung wenden.")
    if "Scripting" in text or "scripting" in text:
        return ("SAP GUI Scripting ist nicht freigeschaltet. In SAP Logon "
                "unter Optionen → Barrierefreiheit & Skripting aktivieren.")
    if "not a zip file" in text.lower() or "BadZipFile" in text:
        return ("Das heruntergeladene Paket ist unvollständig. Bitte den "
                "Lauf für diese Materialnummer wiederholen.")
    return text


class WorkerBridge(QObject):
    """Hebt Orchestrator-Callbacks (Worker-Thread) in den GUI-Thread."""

    progress = Signal(object)
    result = Signal(object)
    log_line = Signal(str)
    finished = Signal(object)

    def callbacks(self) -> Callbacks:
        return Callbacks(
            on_progress=self.progress.emit,
            on_result=self.result.emit,
            on_log=self.log_line.emit,
            on_finished=self.finished.emit,
        )


class MainWindow(QMainWindow):
    def __init__(self, make_adapter, profiles: list[str],
                 mock_default: Path | None = None):
        """make_adapter(config: RunConfig) -> SapAdapter (echtes SAP oder Mock)."""
        super().__init__()
        self.make_adapter = make_adapter
        self.mock_default = mock_default
        self.setWindowTitle("Drawing Checker – Zeichnungsprüfung")
        self.resize(1080, 720)
        self.setAcceptDrops(True)

        self.excel_path: Path | None = None
        self.selected_column: str | None = None
        self.orchestrator: Orchestrator | None = None
        self.adapter: SapAdapter | None = None
        self.run_dir: Path | None = None

        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)
        self.stack.addWidget(self._build_step1(profiles))
        self.stack.addWidget(self._build_step2())
        self.stack.addWidget(self._build_step3())
        self._load_settings()

    # ================================================== Schritt 1: Datei
    def _build_step1(self, profiles: list[str]) -> QWidget:
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.addStretch()
        title = QLabel("<h2>Schritt 1 von 3 – Excel-Datei wählen</h2>")
        title.setAlignment(Qt.AlignCenter)
        lay.addWidget(title)
        hint = QLabel("Die Materialliste Ihrer Materialgruppe (.xlsx) hierher "
                      "ziehen oder über den Button auswählen.")
        hint.setAlignment(Qt.AlignCenter)
        lay.addWidget(hint)

        btn = QPushButton("Datei auswählen …")
        btn.setFixedWidth(220)
        btn.clicked.connect(self._pick_file)
        row = QHBoxLayout()
        row.addStretch(); row.addWidget(btn); row.addStretch()
        lay.addLayout(row)

        self.lbl_file = QLabel("")
        self.lbl_file.setAlignment(Qt.AlignCenter)
        lay.addWidget(self.lbl_file)

        form = QHBoxLayout()
        form.addStretch()
        form.addWidget(QLabel("Tabellenblatt:"))
        self.cmb_sheet = QComboBox()
        self.cmb_sheet.setMinimumWidth(180)
        form.addWidget(self.cmb_sheet)
        form.addSpacing(20)
        form.addWidget(QLabel("Materialgruppe (Regelprofil):"))
        self.cmb_profile = QComboBox()
        self.cmb_profile.addItems(profiles)
        form.addWidget(self.cmb_profile)
        form.addStretch()
        lay.addLayout(form)

        opts = QHBoxLayout()
        opts.addStretch()
        self.chk_resume = QCheckBox("Letzten Lauf fortsetzen")
        opts.addWidget(self.chk_resume)
        opts.addSpacing(20)
        self.chk_mock = QCheckBox("Testmodus (Mockdaten statt SAP)")
        self.chk_mock.setChecked(self.mock_default is not None)
        opts.addWidget(self.chk_mock)
        opts.addSpacing(20)
        opts.addWidget(QLabel("SAP-System:"))
        self.txt_system = QLineEdit("P11")
        self.txt_system.setFixedWidth(70)
        opts.addWidget(self.txt_system)
        opts.addSpacing(12)
        btn_test = QPushButton("Verbindung testen")
        btn_test.clicked.connect(self._test_connection)
        opts.addWidget(btn_test)
        opts.addStretch()
        lay.addLayout(opts)

        nxt = QPushButton("Weiter  ➜")
        nxt.setFixedWidth(160)
        nxt.clicked.connect(self._goto_step2)
        row2 = QHBoxLayout()
        row2.addStretch(); row2.addWidget(nxt); row2.addStretch()
        lay.addLayout(row2)
        lay.addStretch()
        return w

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        for url in e.mimeData().urls():
            p = Path(url.toLocalFile())
            if p.suffix.lower() in (".xlsx", ".xlsm"):
                self._set_file(p)
                return

    def _pick_file(self):
        name, _ = QFileDialog.getOpenFileName(
            self, "Materialliste wählen", "", "Excel-Dateien (*.xlsx *.xlsm)")
        if name:
            self._set_file(Path(name))

    def _set_file(self, path: Path):
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
            sheets = wb.sheetnames
            wb.close()
        except Exception as exc:
            QMessageBox.warning(self, "Datei nicht lesbar",
                                "Die Datei konnte nicht geöffnet werden:\n"
                                + klartext(exc))
            return
        self.excel_path = path
        self.lbl_file.setText(f"<b>{path.name}</b>")
        self.cmb_sheet.clear()
        self.cmb_sheet.addItems(sheets)

    def _test_connection(self):
        """Verbindungstest ohne Prüflauf – für den Durchstich mit SAP."""
        if self.excel_path is None:
            # Config braucht einen Pfad; für den reinen Test genügt ein Dummy.
            self.excel_path = Path("verbindungstest.xlsx")
            dummy = True
        else:
            dummy = False
        try:
            cfg = self._make_config(preview=True)
            adapter = self.make_adapter(cfg)
            adapter.ensure_ready()
            adapter.close()
            QMessageBox.information(
                self, "Verbindungstest",
                ("Mockmodus bereit (Testdatenordner gefunden)."
                 if cfg.mock_source else
                 f"Verbindung zu {cfg.sap_connection} steht – Session "
                 f"gefunden bzw. Anmeldung erfolgreich."))
        except Exception as exc:
            QMessageBox.critical(
                self, "Verbindungstest",
                f"Verbindung nicht möglich:\n{exc}\n\nSAP Logon prüfen "
                "(läuft es? Scripting aktiviert?) und erneut testen.")
        finally:
            if dummy:
                self.excel_path = None

    def _goto_step2(self):
        if self.excel_path is None:
            QMessageBox.information(self, "Datei fehlt",
                                    "Bitte zuerst eine Excel-Datei auswählen.")
            return
        self._fill_preview()
        self.stack.setCurrentIndex(1)

    # ============================================== Schritt 2: Spaltenwahl
    def _build_step2(self) -> QWidget:
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.addWidget(QLabel(
            "<h2>Schritt 2 von 3 – Spalte mit den Materialnummern anklicken</h2>"))
        self.lbl_colinfo = QLabel(
            "Klicken Sie auf die Spaltenüberschrift der Materialnummern.")
        lay.addWidget(self.lbl_colinfo)

        self.table = QTableWidget()
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectColumns)
        self.table.setSelectionMode(QTableWidget.SingleSelection)
        self.table.horizontalHeader().sectionClicked.connect(self._column_clicked)
        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeToContents)
        lay.addWidget(self.table, stretch=1)

        row = QHBoxLayout()
        back = QPushButton("⬅  Zurück")
        back.clicked.connect(lambda: self.stack.setCurrentIndex(0))
        row.addWidget(back)
        row.addStretch()
        self.btn_start = QPushButton("Prüfung starten  ➜")
        self.btn_start.setEnabled(False)
        self.btn_start.setStyleSheet("font-weight: bold; padding: 6px 18px;")
        self.btn_start.clicked.connect(self._start_run)
        row.addWidget(self.btn_start)
        lay.addLayout(row)
        return w

    def _fill_preview(self):
        wb = openpyxl.load_workbook(self.excel_path, read_only=True,
                                    data_only=True)
        ws = wb[self.cmb_sheet.currentText()]
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if r >= PREVIEW_ROWS:
                break
        wb.close()
        ncols = max((len(r) for r in rows), default=0)
        self.table.clear()
        self.table.setRowCount(len(rows))
        self.table.setColumnCount(ncols)
        self.table.setHorizontalHeaderLabels(
            [get_column_letter(c + 1) for c in range(ncols)])
        for r, row in enumerate(rows):
            for c in range(ncols):
                v = row[c] if c < len(row) else None
                self.table.setItem(
                    r, c, QTableWidgetItem("" if v is None else str(v)))
        self.selected_column = None
        self.btn_start.setEnabled(False)
        self._suggest_column(rows)

    def _suggest_column(self, rows) -> None:
        """Schlägt die wahrscheinlichste Materialnummern-Spalte vor.

        Heuristik: Spalte, in der die meisten Zellen wie Materialnummern
        aussehen (6–10 Ziffern). Der Anwender kann jederzeit umklicken.
        """
        import re

        best, best_hits = None, 0
        ncols = self.table.columnCount()
        for c in range(ncols):
            hits = 0
            for row in rows[1:]:
                v = row[c] if c < len(row) else None
                if v is None:
                    continue
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                if re.fullmatch(r"\d{6,10}", str(v).strip()):
                    hits += 1
            if hits > best_hits:
                best, best_hits = c, hits
        if best is not None and best_hits >= 3:
            self._column_clicked(best)
            self.lbl_colinfo.setText(
                self.lbl_colinfo.text()
                + "   (automatisch vorgeschlagen – bei Bedarf andere "
                  "Spalte anklicken)")

    def _column_clicked(self, index: int):
        col = get_column_letter(index + 1)
        self.selected_column = col
        self.table.selectColumn(index)
        # Sofort-Validierung: Wie viele Materialnummern stecken in der Spalte?
        cfg = self._make_config(preview=True)
        try:
            materials = read_materials(cfg)
        except Exception as exc:
            self.lbl_colinfo.setText(f"Spalte {col}: nicht lesbar ({exc})")
            return
        n = len(materials)
        sample = ", ".join(m for _, m in materials[:3])
        self.lbl_colinfo.setText(
            f"Spalte <b>{col}</b> gewählt: <b>{n}</b> Materialnummern erkannt"
            + (f" (z. B. {sample} …)" if sample else "")
            + ("" if n else " – bitte andere Spalte wählen"))
        self.btn_start.setEnabled(n > 0)

    def _make_config(self, preview: bool = False) -> RunConfig:
        assert self.excel_path is not None
        header_row = 1  # Überschrift in Zeile 1; Zeile mit Klick wäre Ausbau
        mock = (self.mock_default
                if (self.chk_mock.isChecked() and self.mock_default) else None)
        return RunConfig(
            excel_path=self.excel_path,
            sheet_name=self.cmb_sheet.currentText(),
            material_column=self.selected_column or "A",
            header_row=header_row,
            output_dir=self.excel_path.parent / "Ergebnisse",
            material_group=self.cmb_profile.currentText(),
            sap_connection=self.txt_system.text().strip() or "P11",
            mock_source=mock,
        )

    # ================================================= Schritt 3: Lauf
    def _build_step3(self) -> QWidget:
        w = QWidget()
        lay = QVBoxLayout(w)
        lay.addWidget(QLabel("<h2>Schritt 3 von 3 – Prüfung läuft</h2>"))
        self.progress_bar = QProgressBar()
        self.progress_bar.setFormat("%v von %m geprüft")
        lay.addWidget(self.progress_bar)
        self.lbl_current = QLabel("")
        lay.addWidget(self.lbl_current)

        splitter = QSplitter(Qt.Horizontal)
        self.result_list = QListWidget()
        self.result_list.currentItemChanged.connect(self._show_detail)
        self.result_list.itemDoubleClicked.connect(self._open_detail_image)
        splitter.addWidget(self.result_list)

        tabs = QTabWidget()
        detail = QWidget()
        dlay = QVBoxLayout(detail)
        self.detail_text = QTextEdit()
        self.detail_text.setReadOnly(True)
        self.detail_text.setPlaceholderText(
            "Ergebnis links anklicken, um Mängel und Zeichnung zu sehen.")
        dlay.addWidget(self.detail_text, stretch=2)
        self.preview = QLabel()
        self.preview.setAlignment(Qt.AlignCenter)
        self.preview.setMinimumHeight(220)
        self.preview.setStyleSheet("background:#e8e8e8; border:1px solid #bbb;")
        dlay.addWidget(self.preview, stretch=3)
        self.btn_image = QPushButton("Zeichnung in voller Größe öffnen")
        self.btn_image.setEnabled(False)
        self.btn_image.clicked.connect(self._open_detail_image)
        dlay.addWidget(self.btn_image)
        tabs.addTab(detail, "Details")
        self.log_view = QTextEdit()
        self.log_view.setReadOnly(True)
        tabs.addTab(self.log_view, "Protokoll")
        splitter.addWidget(tabs)
        splitter.setSizes([340, 720])
        lay.addWidget(splitter, stretch=1)

        row = QHBoxLayout()
        self.btn_pause = QPushButton("Pause")
        self.btn_pause.clicked.connect(self._toggle_pause)
        row.addWidget(self.btn_pause)
        self.btn_stop = QPushButton("Abbrechen")
        self.btn_stop.clicked.connect(self._stop_run)
        row.addWidget(self.btn_stop)
        row.addStretch()
        self.btn_retry = QPushButton("Fehlgeschlagene erneut prüfen")
        self.btn_retry.clicked.connect(self._retry_failed)
        self.btn_retry.setEnabled(False)
        row.addWidget(self.btn_retry)
        self.btn_report = QPushButton("Bericht öffnen")
        self.btn_report.clicked.connect(self._open_report)
        self.btn_report.setEnabled(False)
        row.addWidget(self.btn_report)
        self.btn_open = QPushButton("Ergebnisordner öffnen")
        self.btn_open.clicked.connect(self._open_results)
        self.btn_open.setEnabled(False)
        row.addWidget(self.btn_open)
        self.btn_new = QPushButton("Neue Prüfung")
        self.btn_new.clicked.connect(self._reset)
        self.btn_new.setEnabled(False)
        row.addWidget(self.btn_new)
        lay.addLayout(row)
        return w

    # ------------------------------------------------------ Detailansicht
    SEV_HTML = {
        Severity.INFO: "#4682b4",
        Severity.WARNING: "#e69100",
        Severity.ERROR: "#c81e1e",
        Severity.BLOCKER: "#8c008c",
    }

    def _show_detail(self, item: QListWidgetItem | None, _prev=None) -> None:
        self.preview.clear()
        self.btn_image.setEnabled(False)
        if item is None:
            self.detail_text.clear()
            return
        r: MaterialResult = item.data(Qt.UserRole)
        if r is None:
            return
        parts = [f"<h3>{r.material}</h3>",
                 f"<p>Geprüft am {r.checked_at}"
                 + (f" · letzte Zeichnungsänderung {r.drawing_rev_date}"
                    if r.drawing_rev_date else "")
                 + (f"<br>Fertigungsverfahren: {', '.join(r.processes)}"
                    if r.processes else "") + "</p>"]
        if r.error:
            parts.append(f'<p style="color:#c81e1e"><b>Technischer Fehler:</b> '
                         f"{r.error}</p>")
        if r.step_summary:
            parts.append(f'<p style="color:#555">{r.step_summary}</p>')
        if not r.findings:
            parts.append('<p style="color:#2f7d32"><b>Keine Beanstandungen.'
                         "</b></p>")
        for i, f in enumerate(r.sorted_findings(), start=1):
            color = self.SEV_HTML[f.severity]
            detail = f"<br><small>{f.detail}</small>" if f.detail else ""
            parts.append(
                f'<p style="color:{color}"><b>{i}. [{SEVERITY_LABEL[f.severity]}]'
                f" {f.code}</b><br>{f.text}{detail}</p>")
        self.detail_text.setHtml("".join(parts))

        if r.screenshot and Path(r.screenshot).exists():
            pix = QPixmap(str(r.screenshot))
            if not pix.isNull():
                self.preview.setPixmap(pix.scaled(
                    self.preview.size(), Qt.KeepAspectRatio,
                    Qt.SmoothTransformation))
                self.btn_image.setEnabled(True)

    def _open_detail_image(self, *_):
        item = self.result_list.currentItem()
        r = item.data(Qt.UserRole) if item else None
        if r and r.screenshot and Path(r.screenshot).exists():
            self._open_path(Path(r.screenshot))

    def _open_report(self):
        if self.orchestrator and self.orchestrator.report_path:
            self._open_path(self.orchestrator.report_path)

    def _retry_failed(self):
        self._start_run(resume=True)

    def _start_run(self, resume: bool | None = None):
        if resume is None:
            resume = self.chk_resume.isChecked()
        cfg = self._make_config()
        self._save_settings(cfg)
        try:
            self.adapter = self.make_adapter(cfg)
            self.adapter.ensure_ready()
        except Exception as exc:
            QMessageBox.critical(
                self, "SAP-Verbindung",
                "Die Verbindung konnte nicht hergestellt werden:\n"
                + klartext(exc)
                + "\n\nBitte SAP Logon prüfen und erneut versuchen.")
            return

        self.bridge = WorkerBridge()
        self.bridge.progress.connect(self._on_progress)
        self.bridge.result.connect(self._on_result)
        self.bridge.log_line.connect(self._on_log)
        self.bridge.finished.connect(self._on_finished)
        self.orchestrator = Orchestrator(
            cfg, self.adapter, self.bridge.callbacks(), resume=resume)
        self.run_dir = self.orchestrator.run_dir

        self.result_list.clear()
        self.detail_text.clear()
        self.preview.clear()
        self.log_view.clear()
        self.progress_bar.setValue(0)
        self.btn_pause.setEnabled(True)
        self.btn_pause.setText("Pause")
        self.btn_stop.setEnabled(True)
        self.btn_new.setEnabled(False)
        self.btn_retry.setEnabled(False)
        self.btn_report.setEnabled(False)
        self.stack.setCurrentIndex(2)
        self.orchestrator.start()

    def _toggle_pause(self):
        if self.orchestrator is None:
            return
        if self.orchestrator.pause_event.is_set():
            self.orchestrator.resume()
            self.btn_pause.setText("Pause")
        else:
            self.orchestrator.pause()
            self.btn_pause.setText("Fortsetzen")

    def _stop_run(self):
        if self.orchestrator is None:
            return
        if QMessageBox.question(
                self, "Abbrechen",
                "Prüfung wirklich abbrechen? Bereits geprüfte Ergebnisse "
                "bleiben erhalten und der Lauf kann später fortgesetzt werden."
        ) == QMessageBox.Yes:
            self.orchestrator.stop()

    def _on_progress(self, p: Progress):
        self.progress_bar.setMaximum(max(p.total, 1))
        self.progress_bar.setValue(p.done)
        cur = f"Aktuell: {p.current}" if p.current else ""
        eta = ""
        if p.eta_s is not None and p.eta_s > 5:
            minutes = p.eta_s / 60
            eta = (f"   |   Rest ca. {minutes:.0f} min" if minutes >= 1
                   else "   |   Rest unter 1 min")
        self.lbl_current.setText(
            f"{cur}   |   ✔ {p.ok} ok   ⚠ {p.findings} mit Findings   "
            f"✖ {p.failed} fehlgeschlagen{eta}")

    def _on_result(self, r: MaterialResult):
        texts = {
            JobStatus.OK: "OK",
            JobStatus.FINDINGS: f"{len(r.findings)} Finding(s)",
            JobStatus.FAILED: f"fehlgeschlagen: {r.error}",
            JobStatus.SKIPPED: "übersprungen",
        }
        item = QListWidgetItem(f"{r.material}  –  {texts.get(r.status, '?')}")
        item.setForeground(STATUS_COLOR.get(r.status, QColor(0, 0, 0)))
        item.setData(Qt.UserRole, r)
        self.result_list.addItem(item)
        self.result_list.scrollToBottom()

    def _on_log(self, line: str):
        self.log_view.append(line)

    def _on_finished(self, p: Progress):
        self.btn_pause.setEnabled(False)
        self.btn_stop.setEnabled(False)
        self.btn_open.setEnabled(True)
        self.btn_new.setEnabled(True)
        self.btn_retry.setEnabled(p.failed > 0)
        self.btn_report.setEnabled(
            bool(self.orchestrator and self.orchestrator.report_path))
        self.lbl_current.setText(p.message or "Fertig.")
        QMessageBox.information(self, "Prüfung abgeschlossen",
                                p.message or "Die Prüfung ist abgeschlossen.")

    def _open_results(self):
        target = self.run_dir or (self.excel_path and self.excel_path.parent)
        if target:
            self._open_path(Path(target))

    @staticmethod
    def _open_path(path: Path) -> None:
        if sys.platform == "win32":
            subprocess.Popen(["explorer" if path.is_dir() else "cmd",
                              *([] if path.is_dir() else ["/c", "start", ""]),
                              str(path)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])

    # ---------------------------------------------- Einstellungen merken
    def _save_settings(self, cfg: RunConfig) -> None:
        s = QSettings("DrawingChecker", "DrawingChecker")
        s.setValue("excel_path", str(cfg.excel_path))
        s.setValue("sheet", cfg.sheet_name)
        s.setValue("column", cfg.material_column)
        s.setValue("profile", cfg.material_group)
        s.setValue("system", cfg.sap_connection)

    def _load_settings(self) -> None:
        s = QSettings("DrawingChecker", "DrawingChecker")
        last = s.value("excel_path", "")
        if last and Path(last).exists():
            self._set_file(Path(last))
            sheet = s.value("sheet", "")
            if sheet and self.cmb_sheet.findText(sheet) >= 0:
                self.cmb_sheet.setCurrentText(sheet)
        profile = s.value("profile", "")
        if profile and self.cmb_profile.findText(profile) >= 0:
            self.cmb_profile.setCurrentText(profile)
        self.txt_system.setText(s.value("system", "P11") or "P11")

    def _reset(self):
        self.orchestrator = None
        self.stack.setCurrentIndex(0)


def list_profiles() -> list[str]:
    from ..checks.base import load_profiles_data

    names = list(load_profiles_data())
    names.sort(key=lambda n: (n != "default", n))
    return names
