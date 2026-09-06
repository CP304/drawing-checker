"""Excel-Rückschrieb: Ergebnis-Spalten in eine Kopie der Input-Datei.

Das Original bleibt unangetastet; geschrieben wird `<name>_geprüft.xlsx`.
Nach jeder Materialnummer wird gespeichert, damit auch bei Abbruch ein
verwertbarer Stand existiert.
"""
from __future__ import annotations

import logging
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from ..core.models import JobStatus, MaterialResult, RunConfig, Severity, SEVERITY_LABEL

log = logging.getLogger(__name__)

RESULT_HEADERS = [
    "Prüfstatus", "Schwerste Bewertung", "Anzahl Findings",
    "Findings", "Geometrieabgleich", "Screenshot", "Prüfdauer [s]",
]
FILL = {
    "ok": PatternFill("solid", fgColor="C6EFCE"),
    "warn": PatternFill("solid", fgColor="FFEB9C"),
    "error": PatternFill("solid", fgColor="FFC7CE"),
    "failed": PatternFill("solid", fgColor="D9D9D9"),
}
STATUS_TEXT = {
    JobStatus.OK: "OK",
    JobStatus.FINDINGS: "Findings",
    JobStatus.FAILED: "Prüfung fehlgeschlagen",
    JobStatus.SKIPPED: "Übersprungen",
}


class ResultWorkbook:
    def __init__(self, config: RunConfig):
        self.config = config
        self.path = config.excel_path.with_stem(config.excel_path.stem + "_geprüft")
        if not self.path.exists():
            shutil.copy2(config.excel_path, self.path)
        self.wb = openpyxl.load_workbook(self.path)
        if config.sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Blatt {config.sheet_name!r} nicht in {self.path.name}")
        self.ws = self.wb[config.sheet_name]
        self.first_col = self._ensure_headers()

    def _ensure_headers(self) -> int:
        """Findet oder erzeugt die Ergebnis-Spalten rechts der Tabelle."""
        header_row = self.config.header_row
        existing = {
            (c.value or ""): c.column
            for c in self.ws[header_row]
            if isinstance(c.value, str)
        }
        if RESULT_HEADERS[0] in existing:
            return existing[RESULT_HEADERS[0]]
        first = (self.ws.max_column or 0) + 1
        for i, name in enumerate(RESULT_HEADERS):
            cell = self.ws.cell(row=header_row, column=first + i, value=name)
            cell.font = Font(bold=True)
            self.ws.column_dimensions[get_column_letter(first + i)].width = (
                46 if name in ("Findings", "Geometrieabgleich") else 18)
        return first

    def write_result(self, result: MaterialResult) -> None:
        r = result.row
        c = self.first_col
        worst = result.worst_severity
        findings_txt = "\n".join(
            f"[{SEVERITY_LABEL[f.severity]}] {f.code}: {f.text}"
            + (f" – {f.detail}" if f.detail else "")
            for f in result.sorted_findings()
        )
        status_txt = STATUS_TEXT.get(result.status, result.status.value)
        if result.status == JobStatus.FAILED and result.error:
            status_txt += f": {result.error}"

        values = [
            status_txt,
            SEVERITY_LABEL[worst] if worst is not None else "",
            len(result.findings),
            findings_txt,
            result.step_summary,
            "",  # Screenshot als Hyperlink, s. u.
            round(result.duration_s, 1),
        ]
        for i, v in enumerate(values):
            cell = self.ws.cell(row=r, column=c + i, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        if result.screenshot:
            cell = self.ws.cell(row=r, column=c + 5)
            cell.value = result.screenshot.name
            cell.hyperlink = result.screenshot.resolve().as_uri()
            cell.font = Font(color="0563C1", underline="single")

        status_cell = self.ws.cell(row=r, column=c)
        if result.status == JobStatus.FAILED:
            status_cell.fill = FILL["failed"]
        elif worst is None or worst <= Severity.INFO:
            status_cell.fill = FILL["ok"]
        elif worst == Severity.WARNING:
            status_cell.fill = FILL["warn"]
        else:
            status_cell.fill = FILL["error"]

    def save(self) -> None:
        try:
            self.wb.save(self.path)
        except PermissionError:
            # Datei ist vermutlich in Excel geöffnet – Ausweichdatei schreiben.
            alt = self.path.with_stem(self.path.stem + "_neu")
            log.warning("Ergebnisdatei gesperrt, schreibe nach %s", alt.name)
            self.wb.save(alt)


def read_materials(config: RunConfig) -> list[tuple[int, str]]:
    """Liest (Zeile, Materialnummer) aus der gewählten Spalte der Input-Excel."""
    wb = openpyxl.load_workbook(config.excel_path, read_only=True, data_only=True)
    try:
        ws = wb[config.sheet_name]
        col = column_index_from_string(config.material_column)
        out: list[tuple[int, str]] = []
        for row in ws.iter_rows(min_row=config.header_row + 1,
                                min_col=col, max_col=col):
            cell = row[0]
            value = cell.value
            if value is None:
                continue
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            text = str(value).strip()
            if text:
                out.append((cell.row, text))
        return out
    finally:
        wb.close()
