"""Laufabschluss-Artefakte: HTML-Bericht, Excel-Zusammenfassung, Lauf-Log."""
from pathlib import Path

import openpyxl
import pytest

from drawing_checker.kern import RunConfig
from drawing_checker.ablauf import Callbacks, Orchestrator
from drawing_checker.sap_ymatdocs import MockSapAdapter


@pytest.fixture(scope="module")
def finished_run(mock_dir, tmp_path_factory):
    out = tmp_path_factory.mktemp("erg")
    cfg = RunConfig(
        excel_path=mock_dir / "Materialliste_Mock.xlsx",
        sheet_name="Materialliste", material_column="C", header_row=1,
        output_dir=out)
    adapter = MockSapAdapter(mock_dir)
    adapter.ensure_ready()
    orch = Orchestrator(cfg, adapter, Callbacks())
    orch.start(); orch.join(300)
    return orch


def test_html_report_written(finished_run):
    path = finished_run.report_path
    assert path is not None and path.exists()
    html = path.read_text(encoding="utf-8")
    assert "Zeichnungsprüfung – Bericht" in html
    assert "10473216" in html          # Geometrie-K.O.-Fall
    assert "GEO.MISMATCH" in html
    assert "10473216.png" in html      # Link auf annotierte Zeichnung


def test_excel_summary_sheet(finished_run, mock_dir):
    wb = openpyxl.load_workbook(mock_dir / "Materialliste_Mock_geprüft.xlsx")
    assert "Prüfzusammenfassung" in wb.sheetnames
    ws = wb["Prüfzusammenfassung"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "LANG.GERMAN" in values     # häufigster Mängelcode gezählt
    # Ergebnisblatt hat Autofilter und Fixierung
    main = wb["Materialliste"]
    assert main.auto_filter.ref
    assert main.freeze_panes == "A2"


def test_run_logfile_written(finished_run):
    assert (finished_run.run_dir / "lauf.log").read_text(encoding="utf-8")


def test_annotated_image_has_stamp(finished_run):
    # Stempel oben links: Pixel im Rahmenbereich sind nicht mehr weiß
    from PIL import Image

    shot = next(r.screenshot for r in finished_run.state.results.values()
                if r.material == "10473216")
    img = Image.open(shot).convert("RGB")
    box = img.crop((16, 16, 200, 70))
    assert any(p != (255, 255, 255) for p in box.getdata())
