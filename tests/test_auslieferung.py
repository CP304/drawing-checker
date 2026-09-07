"""Auslieferung: Paketbau, Einzeldatei, Startskript, Doku, Modulhygiene.
"""
from __future__ import annotations

# ======================================================================
# paket
# ======================================================================
# Prüft das Verteilpaket: eine ZIP-Datei, die am Zielrechner reicht.
#
# Der Anwenderrechner bekommt genau eine Datei. Fehlt darin ein
# Wissenspaket oder das Startskript, merkt man es erst dort – deshalb wird
# das Archiv hier gebaut und gegengeprüft.
import zipfile
from pathlib import Path

import pytest

from tools import paket

WURZEL = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def archiv(tmp_path_factory) -> Path:
    ziel = tmp_path_factory.mktemp("paket")
    return paket.zip_bauen(ziel, mit_tests=False)


def namen(archiv: Path) -> set[str]:
    with zipfile.ZipFile(archiv) as zf:
        return set(zf.namelist())


def test_alles_liegt_in_einem_ordner(archiv):
    """Beim Entpacken darf nichts verstreut im Zielordner landen."""
    assert all(n.startswith("DrawingChecker/") for n in namen(archiv))


@pytest.mark.parametrize("datei", [
    "Start.bat", "LIESMICH.txt", "pyproject.toml", "README.md",
    "drawing_checker/app.py",
    "drawing_checker/gui.py",
    "drawing_checker/sap_ymatdocs.py",
    "drawing_checker/rules/profiles.yaml",
    "drawing_checker/rules/materials.yaml",
    "drawing_checker/rules/norms.yaml",
    "drawing_checker/rules/beschaffung.yaml",
])
def test_pflichtdateien_enthalten(archiv, datei):
    assert f"DrawingChecker/{datei}" in namen(archiv)


def test_ballast_bleibt_draussen(archiv):
    """Kalibrierzeichnungen (25 MB) und Caches gehören nicht ins Paket."""
    for n in namen(archiv):
        assert "echt_quellen" not in n
        assert "__pycache__" not in n and not n.endswith(".pyc")
        assert "/.venv/" not in n
    assert archiv.stat().st_size < 5 * 1024 * 1024, "Paket unerwartet groß"


def test_archiv_ist_lesbar_und_vollstaendig(archiv):
    """Baut, entpackt und lädt die Wissenspakete im entpackten Stand."""
    assert paket.zip_pruefen(archiv) == []


def test_mit_tests_enthaelt_die_testsuite(tmp_path):
    voll = paket.zip_bauen(tmp_path, mit_tests=True)
    inhalt = namen(voll)
    assert "DrawingChecker/tests/test_ablauf.py" in inhalt
    assert "DrawingChecker/mockdata/daten.py" in inhalt


# ------------------------------------------------- Einzeldatei (self-extract)
def test_einzeldatei_enthaelt_ein_gueltiges_paket(tmp_path):
    """Die selbstentpackende .bat muss ein brauchbares ZIP tragen."""
    
    datei = paket.bat_bauen(tmp_path)
    assert datei.name == "DrawingChecker_Setup.bat"
    assert paket.bat_pruefen(datei) == []


def test_einzeldatei_hat_marke_und_windows_zeilenenden(tmp_path):
    
    datei = paket.bat_bauen(tmp_path)
    roh = datei.read_bytes()
    assert b"::PAYLOAD::\r\n" in roh
    assert roh.startswith(b"@echo off\r\n")
    # Der Batch-Teil darf die Nutzlast nie ausfuehren. Getrennt wird an der
    # Markenzeile, nicht am ersten Vorkommen - die Marke steht auch im
    # PowerShell-Aufruf des Kopfes.
    kopf = roh.split(b"\r\n::PAYLOAD::\r\n")[0].decode("ascii")
    assert "exit /b 0" in kopf
    assert kopf.count("::PAYLOAD::") >= 1        # Suche im Kopf vorhanden


def test_einzeldatei_nutzlast_ist_das_paket(tmp_path):
    """Was drinsteckt, ist genau das gebaute ZIP."""
    import zipfile

    
    zip_pfad = paket.zip_bauen(tmp_path, mit_tests=False)
    datei = paket.bat_bauen(tmp_path, quelle=zip_pfad)
    assert paket.nutzlast(datei) == zip_pfad.read_bytes()
    with zipfile.ZipFile(zip_pfad) as zf:
        assert "DrawingChecker/Start.bat" in zf.namelist()


# ------------------------------------------- Kalibrierzeichnungen als Archiv
def test_kalibrierzeichnungen_liegen_als_ein_archiv():
    """Über hundert Einzeldateien im Repo waren unübersichtlich."""
    from mockdata import daten as quellen

    assert quellen.ARCHIV.is_file()
    pdf, step = quellen.anzahl()
    assert pdf >= 80 and step >= 25


def test_zeichnungen_werden_ausgepackt(tmp_path):
    from mockdata.daten import zeichnungen

    ordner = zeichnungen(ziel=tmp_path / "raus")
    assert len(list(ordner.glob("*.pdf"))) >= 80
    # Zweiter Aufruf packt nicht erneut aus.
    marke = (ordner / ".ausgepackt").stat().st_mtime_ns
    zeichnungen(ziel=tmp_path / "raus")
    assert (ordner / ".ausgepackt").stat().st_mtime_ns == marke


# ======================================================================
# package
# ======================================================================


from drawing_checker.kern import ( PackageError, classify_files, extract_package, )


def make_zip(path: Path, names: dict[str, bytes]):
    with zipfile.ZipFile(path, "w") as zf:
        for name, data in names.items():
            zf.writestr(name, data)


def test_extract_classifies_content(tmp_path):
    z = tmp_path / "m1.zip"
    make_zip(z, {
        "Z_123.pdf": b"%PDF-1.4 x",
        "M_123.stp": b"ISO-10303-21;",
        "N_123.CATPart": b"native",
    })
    c = extract_package(z, tmp_path / "work", "123")
    assert c.drawing_pdf.name == "Z_123.pdf"
    assert c.step_file.name == "M_123.stp"
    assert len(c.ignored) == 1


def test_zip_slip_names_flattened(tmp_path):
    z = tmp_path / "m2.zip"
    make_zip(z, {"../../evil.pdf": b"x", "sub/dir/ok.pdf": b"y"})
    c = extract_package(z, tmp_path / "work", "m2")
    names = {p.name for p in c.pdfs}
    assert names == {"evil.pdf", "ok.pdf"}
    for p in c.pdfs:
        assert (tmp_path / "work") in p.parents


def test_multi_pdf_prefers_material_in_name(tmp_path):
    a = tmp_path / "Anbau.pdf"; a.write_bytes(b"x" * 500)
    b = tmp_path / "Z_10473215.pdf"; b.write_bytes(b"x" * 100)
    c = classify_files([a, b], "10473215")
    assert c.drawing_pdf.name == "Z_10473215.pdf"


def test_missing_zip_raises(tmp_path):
    with pytest.raises(PackageError):
        extract_package(tmp_path / "fehlt.zip", tmp_path / "w", "x")


def test_empty_zip_raises(tmp_path):
    z = tmp_path / "leer.zip"
    z.write_bytes(b"")
    with pytest.raises(PackageError):
        extract_package(z, tmp_path / "w", "x")


# ------------------------------------------------- Schranke fuers Zusammenlegen
def test_kein_name_wird_im_modul_doppelt_vergeben():
    """Zusammengelegte Module duerfen sich nicht gegenseitig ueberschreiben.

    Beim Flachziehen der Paketstruktur sind zwei verschiedene Regexe unter
    demselben Namen `RE_SCALE` in einem Modul gelandet - der zweite hat den
    ersten verdeckt und die Maßstabserkennung stillgelegt. Gefunden haben
    das die Tests; damit es gar nicht erst passiert, prueft dieser Test
    jedes Modul auf doppelt vergebene Namen auf oberster Ebene.
    """
    import ast
    from collections import Counter
    from pathlib import Path

    wurzel = Path(__file__).resolve().parent.parent / "drawing_checker"
    doppelt = {}
    for pfad in sorted(wurzel.glob("*.py")):
        namen = Counter()
        for knoten in ast.parse(pfad.read_text(encoding="utf-8")).body:
            if isinstance(knoten, (ast.FunctionDef, ast.AsyncFunctionDef,
                                   ast.ClassDef)):
                namen[knoten.name] += 1
            elif isinstance(knoten, ast.Assign):
                for ziel in knoten.targets:
                    if isinstance(ziel, ast.Name):
                        namen[ziel.id] += 1
        mehrfach = {n: z for n, z in namen.items() if z > 1}
        if mehrfach:
            doppelt[pfad.name] = mehrfach
    assert not doppelt, f"Namen doppelt vergeben: {doppelt}"


# ======================================================================
# startskript
# ======================================================================
# Prüft Start.bat auf die klassischen Batch-Fallen.
#
# Unter Linux lässt sich eine .bat-Datei nicht ausführen; die typischen
# Fehler sind aber statisch erkennbar und haben es in sich – eine unescapte
# Klammer in einem `if (...)`-Block bricht die Datei mitten im Lauf ab, und
# das merkt man erst beim Anwender.
import re


BAT = Path(__file__).resolve().parent.parent / "Start.bat"
TEXT = BAT.read_text(encoding="ascii", errors="strict")
ZEILEN = TEXT.splitlines()


def test_datei_vorhanden_und_reines_ascii():
    """Umlaute in .bat-Dateien werden je nach Codepage zu Buchstabensalat."""
    assert BAT.is_file()
    TEXT.encode("ascii")          # wirft bei Umlauten


def _klammer_tiefe(zeile: str, tiefe: int) -> int:
    """Blocktiefe fortschreiben: Anführungszeichen und ^-Escapes beachten."""
    in_string = False
    i = 0
    while i < len(zeile):
        c = zeile[i]
        if c == "^":
            i += 2                # nächstes Zeichen ist escaped
            continue
        if c == '"':
            in_string = not in_string
        elif not in_string and c == "(":
            tiefe += 1
        elif not in_string and c == ")":
            tiefe -= 1
        i += 1
    return tiefe


def test_bloecke_sind_ausgeglichen():
    tiefe = 0
    for nr, zeile in enumerate(ZEILEN, start=1):
        if zeile.strip().lower().startswith("rem "):
            continue
        tiefe = _klammer_tiefe(zeile, tiefe)
        assert tiefe >= 0, f"Zeile {nr}: schließende Klammer zu viel"
    assert tiefe == 0, "am Dateiende ist ein Block noch offen"


def test_alle_sprungziele_existieren():
    labels = {z.strip().lstrip(":").lower()
              for z in ZEILEN if z.strip().startswith(":")}
    ziele = {m.group(1).lower()
             for m in re.finditer(r"goto\s+:?(\w+)", TEXT, re.IGNORECASE)}
    fehlend = ziele - labels
    assert not fehlend, f"Sprungziele ohne Label: {sorted(fehlend)}"


def test_keine_verzoegerte_expansion_noetig():
    """Variablen, die in einem Block gesetzt UND gelesen werden.

    Ohne `setlocal EnableDelayedExpansion` liest %VAR% dort den ALTEN
    Wert – der Klassiker unter den Batch-Fehlern.
    """
    tiefe = 0
    gesetzt_im_block: set[str] = set()
    probleme: list[str] = []
    for nr, zeile in enumerate(ZEILEN, start=1):
        if zeile.strip().lower().startswith("rem "):
            continue
        vorher = tiefe
        tiefe = _klammer_tiefe(zeile, tiefe)
        if vorher == 0 and tiefe > 0:
            gesetzt_im_block = set()
        if tiefe > 0 or vorher > 0:
            for m in re.finditer(r'set\s+"?(\w+)=', zeile, re.IGNORECASE):
                gesetzt_im_block.add(m.group(1).lower())
            for m in re.finditer(r"%(\w+)%", zeile):
                if m.group(1).lower() in gesetzt_im_block:
                    probleme.append(f"Zeile {nr}: %{m.group(1)}%")
        if tiefe == 0:
            gesetzt_im_block = set()
    assert not probleme, ("verzögerte Expansion nötig oder Zuweisung "
                          f"umstellen: {probleme}")


def test_ohne_argument_wird_nichts_durchgereicht():
    """`Start.bat neu`/`pruefen` dürfen nicht an das Programm gehen."""
    assert 'if /I "%~1"=="neu" (' in TEXT
    assert re.search(r'if /I "%~1"=="neu" \(\s*\n\s*set "REBUILD=1"\s*\n\s*'
                     r'set "ARGS="', TEXT)


@pytest.mark.parametrize("schritt", [
    "-m venv",                       # Umgebung anlegen
    "-m pip install -e \".\"",       # Programm installieren
    '.[occ]', '.[ocr]', '.[sap]',    # Zusatzpakete
    "-m drawing_checker.app --check-rules",
    "-m drawing_checker.app %ARGS%",  # Start mit durchgereichten Optionen
    "-m drawing_checker.app --sap-import-vbs",
    "-m drawing_checker.app --sap-dry-run",
    "-m drawing_checker.app --sap-test",
    "-m drawing_checker.app --sap-dump",
    "-m pytest tests -q",            # Selbsttest
])
def test_wesentliche_schritte_vorhanden(schritt):
    assert schritt in TEXT, f"Schritt fehlt im Startskript: {schritt}"


def _labelblock(label: str) -> str:
    """Text ab der Label-DEFINITION (nicht ab dem goto) bis exit /b."""
    m = re.search(rf"^:{label}\s*$", TEXT, re.MULTILINE)
    assert m, f"Label :{label} fehlt"
    return TEXT[m.end():].split("exit /b", 1)[0]


def test_fehlerwege_halten_das_fenster_offen():
    """Bei Doppelklick darf das Fenster im Fehlerfall nicht zuklappen."""
    for label in ("kein_python", "fehler_venv", "fehler_pip", "fehler_lauf"):
        assert "pause" in _labelblock(label), \
            f"{label}: kein pause vor dem Beenden"


def test_hilfetext_nennt_alle_varianten():
    hilfe = _labelblock("hilfe")
    for variante in ("neu", "pruefen", "--sap-import-vbs"):
        assert variante in hilfe


# --------------------------------------------------------------- Menue
def test_menue_deckt_alle_schritte_von_morgen_ab():
    """Ohne Argumente muss ein Menue kommen - morgen tippt niemand Befehle."""
    menue = _labelblock("menu")
    for eintrag in ("Zeichnungen pruefen", "SAP-Mitschnitt einlesen",
                    "Trockenlauf ohne SAP", "Materialnummer testweise",
                    "SAP-Bild anzeigen", "Installation und Regeln pruefen",
                    "Anleitung oeffnen", "Beenden"):
        assert eintrag in menue, f"Menuepunkt fehlt: {eintrag}"


def test_jede_menuewahl_hat_ein_ziel():
    menue = _labelblock("menu")
    ziele = re.findall(r'if "%WAHL%"=="(\d)" goto :(\w+)', menue)
    assert len(ziele) == 8, f"nicht 8 Menuepunkte verdrahtet: {ziele}"
    labels = {z.strip().lstrip(":").lower()
              for z in ZEILEN if z.strip().startswith(":")}
    for nummer, ziel in ziele:
        assert ziel.lower() in labels, f"Punkt {nummer} zeigt auf :{ziel}"


def test_menue_kehrt_zurueck():
    """Nach jeder Aktion muss man wieder im Menue landen."""
    for label in ("m_start", "m_vbs", "m_trocken", "m_test", "m_dump",
                  "m_anleitung"):
        block = _labelblock(label)
        assert "goto :menu" in block, f"{label} kehrt nicht ins Menue zurueck"


def test_warnt_beim_start_aus_dem_zip():
    """Aus dem ZIP heraus gestartet gingen alle Ergebnisse verloren."""
    assert 'find /I "\\Temp\\"' in TEXT
    block = _labelblock("aus_zip")
    assert "entpacken" in block.lower()


# ======================================================================
# documentation
# ======================================================================
# Prüfdokumentation: Zeitstempel, Änderungsdatum, Fertigungsverfahren, Excel.

import openpyxl
import pymupdf

from drawing_checker.zeichnung import extract_revision_date
from drawing_checker.zeichnung import DrawingPdf
from conftest import FONT



def make_pdf(tmp_path: Path, lines) -> DrawingPdf:
    doc = pymupdf.open()
    page = doc.new_page(width=842, height=595)
    kwargs = {}
    if Path(FONT).exists():
        page.insert_font(fontname="T", fontfile=FONT)
        kwargs["fontname"] = "T"
    lines = list(lines) + ["Interne Testzeichnung – Blatt 1 von 1"]
    for i, line in enumerate(lines):
        page.insert_text((40, 60 + i * 30), line, fontsize=10, **kwargs)
    path = tmp_path / "t.pdf"
    doc.save(path)
    doc.close()
    return DrawingPdf(path)


# ------------------------------------------------------- Änderungsdatum
def test_latest_date_wins(tmp_path):
    pdf = make_pdf(tmp_path, [
        "Erstellt 14.11.2024", "Änderung A 02.03.2025",
        "Änderung B 2026-01-12", "Geprüft 15.11.2024"])
    assert extract_revision_date(pdf) == "2026-01-12"


def test_us_dates_parsed(tmp_path):
    pdf = make_pdf(tmp_path, ["REV B 5/23/2021", "REV C 7/14/2021"])
    assert extract_revision_date(pdf) == "2021-07-14"


def test_norm_years_are_not_dates(tmp_path):
    pdf = make_pdf(tmp_path, ["ISO 2768:1989", "ISO 8062-3:2007",
                              "Stand 03.05.2023"])
    assert extract_revision_date(pdf) == "2023-05-03"


def test_pdf_metadata_fallback(tmp_path):
    pdf = make_pdf(tmp_path, ["keine Datumsangabe im Text"])
    result = extract_revision_date(pdf)
    # PyMuPDF setzt kein CreationDate -> leer ist hier korrekt; Hauptsache
    # kein Absturz und kein erfundenes Datum.
    assert result == "" or "PDF-Metadatum" in result


# -------------------------------------------------- Fertigungsverfahren
def test_detect_processes(tmp_path):
    from drawing_checker.pruef_zeichnung import detect_processes

    pdf = make_pdf(tmp_path, [
        "Werkstoff EN-GJS-400-15, Gussteil nach ISO 8062",
        "Lagerbohrung ⌀90 H7, Ra 6,3",
        "Gewinde M12", "lackiert RAL 7016",
    ])
    procs = detect_processes(pdf)
    assert "Gießen" in procs
    assert "Spanende Bearbeitung" in procs
    assert "Gewindefertigung" in procs
    assert "Lackieren/Beschichten" in procs
    assert "Schweißen" not in procs


def test_detect_processes_from_castable_material_only(tmp_path):
    from drawing_checker.pruef_zeichnung import detect_processes

    pdf = make_pdf(tmp_path, ["Werkstoff EN-GJL-250"])
    assert "Gießen" in detect_processes(pdf)


# --------------------------------------------------- Excel-Dokumentation
def test_excel_contains_documentation_columns(mock_dir, tmp_path):
    from drawing_checker.kern import RunConfig
    from drawing_checker.ablauf import Callbacks, Orchestrator
    from drawing_checker.bericht import RESULT_HEADERS
    from drawing_checker.sap_ymatdocs import MockSapAdapter

    cfg = RunConfig(
        excel_path=mock_dir / "Materialliste_Mock.xlsx",
        sheet_name="Materialliste", material_column="C", header_row=1,
        output_dir=tmp_path / "erg")
    adapter = MockSapAdapter(mock_dir)
    adapter.ensure_ready()
    orch = Orchestrator(cfg, adapter, Callbacks())
    orch.start(); orch.join(300)

    ws = openpyxl.load_workbook(
        mock_dir / "Materialliste_Mock_geprüft.xlsx")["Materialliste"]
    headers = {c.value: c.column for c in ws[1] if isinstance(c.value, str)}
    for required in ("Geprüft am", "Screenshot", "Festgestellte Mängel",
                     "Letzte Zeichnungsänderung", "Fertigungsverfahren"):
        assert required in headers, f"Spalte {required} fehlt"

    # Welle (Zeile 4): alle Dokumentationsfelder gefüllt
    row = 4
    assert ws.cell(row=row, column=headers["Geprüft am"]).value
    assert ws.cell(row=row, column=headers["Letzte Zeichnungsänderung"]
                   ).value == "2026-01-12"
    procs = ws.cell(row=row, column=headers["Fertigungsverfahren"]).value
    assert "Spanende Bearbeitung" in procs and "Gewindefertigung" in procs
    shot = ws.cell(row=row, column=headers["Screenshot"])
    assert shot.hyperlink is not None
    # Schweißkonsole (Zeile 2): Schweißen erkannt
    procs2 = ws.cell(row=2, column=headers["Fertigungsverfahren"]).value
    assert "Schweißen" in procs2
