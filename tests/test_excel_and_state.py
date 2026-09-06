from pathlib import Path

import openpyxl

from drawing_checker.core.models import (
    BBox, Finding, JobStatus, MaterialResult, RunConfig, Severity,
)
from drawing_checker.core.state import RunState
from drawing_checker.report.excel_writer import (
    RESULT_HEADERS, ResultWorkbook, read_materials,
)


def make_config(tmp_path: Path, excel: Path) -> RunConfig:
    return RunConfig(excel_path=excel, sheet_name="Materialliste",
                     material_column="C", header_row=1,
                     output_dir=tmp_path / "out")


def test_read_materials_skips_blanks(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materialliste"
    ws.append(["Nr", "Werk", "Materialnummer"])
    ws.append([1, "1000", "10473215"])
    ws.append([2, "1000", None])
    ws.append([3, "1000", 10473216])   # als Zahl formatiert
    ws.append([4, "1000", "  "])
    excel = tmp_path / "liste.xlsx"
    wb.save(excel)

    mats = read_materials(make_config(tmp_path, excel))
    assert mats == [(2, "10473215"), (4, "10473216")]


def test_result_workbook_roundtrip(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materialliste"
    ws.append(["Nr", "Werk", "Materialnummer"])
    ws.append([1, "1000", "10473215"])
    excel = tmp_path / "liste.xlsx"
    wb.save(excel)

    cfg = make_config(tmp_path, excel)
    rwb = ResultWorkbook(cfg)
    result = MaterialResult(material="10473215", row=2,
                            status=JobStatus.FINDINGS)
    result.findings.append(Finding("LANG.GERMAN", Severity.ERROR, "Test",
                                   BBox(0, 0, 1, 1)))
    rwb.write_result(result)
    rwb.save()

    out = openpyxl.load_workbook(rwb.path)["Materialliste"]
    headers = [c.value for c in out[1]]
    assert RESULT_HEADERS[0] in headers
    col = headers.index(RESULT_HEADERS[0]) + 1
    assert out.cell(row=2, column=col).value == "Findings"

    # Zweites Öffnen erzeugt KEINE doppelten Spalten
    rwb2 = ResultWorkbook(cfg)
    assert rwb2.first_col == rwb.first_col


def test_state_roundtrip_and_resume(tmp_path):
    excel = tmp_path / "l.xlsx"
    wb = openpyxl.Workbook(); wb.active.title = "Materialliste"; wb.save(excel)
    cfg = make_config(tmp_path, excel)
    run_dir = tmp_path / "lauf"
    run_dir.mkdir()

    state = RunState(cfg, run_dir)
    r = MaterialResult(material="10473215", row=2, status=JobStatus.OK,
                       duration_s=1.5)
    r.findings.append(Finding("X", Severity.WARNING, "t", BBox(1, 2, 3, 4),
                              page=0, detail="d"))
    state.record(r)

    loaded = RunState.load(cfg, run_dir)
    assert loaded.is_done(2, "10473215")
    assert not loaded.is_done(3, "10473215")
    lr = loaded.results["2:10473215"]
    assert lr.findings[0].severity == Severity.WARNING
    assert lr.findings[0].bbox.x1 == 3


def test_failed_jobs_are_retried_on_resume(tmp_path):
    excel = tmp_path / "l.xlsx"
    wb = openpyxl.Workbook(); wb.active.title = "Materialliste"; wb.save(excel)
    cfg = make_config(tmp_path, excel)
    run_dir = tmp_path / "lauf"; run_dir.mkdir()
    state = RunState(cfg, run_dir)
    state.record(MaterialResult(material="M", row=5, status=JobStatus.FAILED,
                                error="SAP weg"))
    loaded = RunState.load(cfg, run_dir)
    assert not loaded.is_done(5, "M")   # FAILED wird beim Fortsetzen erneut geprüft
