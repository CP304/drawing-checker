"""Prüfdokumentation: Zeitstempel, Änderungsdatum, Fertigungsverfahren, Excel."""
from pathlib import Path

import openpyxl
import pymupdf

from drawing_checker.zeichnung import extract_revision_date
from drawing_checker.zeichnung import DrawingPdf

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def make_pdf(tmp_path: Path, lines) -> DrawingPdf:
    doc = pymupdf.open()
    page = doc.new_page(width=842, height=595)
    kwargs = {}
    if Path(FONT).exists():
        page.insert_font(fontname="T", fontfile=FONT)
        kwargs["fontname"] = "T"
    lines = list(lines) + ["Interne Testzeichnung – Blatt 1 von 1"]
    for i, line in enumerate(lines):
        page.insert_text((40, 60 + i * 30), line, fontsize=10, **kwargs)
    path = tmp_path / "t.pdf"
    doc.save(path)
    doc.close()
    return DrawingPdf(path)


# ------------------------------------------------------- Änderungsdatum
def test_latest_date_wins(tmp_path):
    pdf = make_pdf(tmp_path, [
        "Erstellt 14.11.2024", "Änderung A 02.03.2025",
        "Änderung B 2026-01-12", "Geprüft 15.11.2024"])
    assert extract_revision_date(pdf) == "2026-01-12"


def test_us_dates_parsed(tmp_path):
    pdf = make_pdf(tmp_path, ["REV B 5/23/2021", "REV C 7/14/2021"])
    assert extract_revision_date(pdf) == "2021-07-14"


def test_norm_years_are_not_dates(tmp_path):
    pdf = make_pdf(tmp_path, ["ISO 2768:1989", "ISO 8062-3:2007",
                              "Stand 03.05.2023"])
    assert extract_revision_date(pdf) == "2023-05-03"


def test_pdf_metadata_fallback(tmp_path):
    pdf = make_pdf(tmp_path, ["keine Datumsangabe im Text"])
    result = extract_revision_date(pdf)
    # PyMuPDF setzt kein CreationDate -> leer ist hier korrekt; Hauptsache
    # kein Absturz und kein erfundenes Datum.
    assert result == "" or "PDF-Metadatum" in result


# -------------------------------------------------- Fertigungsverfahren
def test_detect_processes(tmp_path):
    from drawing_checker.pruef_zeichnung import detect_processes

    pdf = make_pdf(tmp_path, [
        "Werkstoff EN-GJS-400-15, Gussteil nach ISO 8062",
        "Lagerbohrung ⌀90 H7, Ra 6,3",
        "Gewinde M12", "lackiert RAL 7016",
    ])
    procs = detect_processes(pdf)
    assert "Gießen" in procs
    assert "Spanende Bearbeitung" in procs
    assert "Gewindefertigung" in procs
    assert "Lackieren/Beschichten" in procs
    assert "Schweißen" not in procs


def test_detect_processes_from_castable_material_only(tmp_path):
    from drawing_checker.pruef_zeichnung import detect_processes

    pdf = make_pdf(tmp_path, ["Werkstoff EN-GJL-250"])
    assert "Gießen" in detect_processes(pdf)


# --------------------------------------------------- Excel-Dokumentation
def test_excel_contains_documentation_columns(mock_dir, tmp_path):
    from drawing_checker.kern import RunConfig
    from drawing_checker.ablauf import Callbacks, Orchestrator
    from drawing_checker.bericht import RESULT_HEADERS
    from drawing_checker.sap_ymatdocs import MockSapAdapter

    cfg = RunConfig(
        excel_path=mock_dir / "Materialliste_Mock.xlsx",
        sheet_name="Materialliste", material_column="C", header_row=1,
        output_dir=tmp_path / "erg")
    adapter = MockSapAdapter(mock_dir)
    adapter.ensure_ready()
    orch = Orchestrator(cfg, adapter, Callbacks())
    orch.start(); orch.join(300)

    ws = openpyxl.load_workbook(
        mock_dir / "Materialliste_Mock_geprüft.xlsx")["Materialliste"]
    headers = {c.value: c.column for c in ws[1] if isinstance(c.value, str)}
    for required in ("Geprüft am", "Screenshot", "Festgestellte Mängel",
                     "Letzte Zeichnungsänderung", "Fertigungsverfahren"):
        assert required in headers, f"Spalte {required} fehlt"

    # Welle (Zeile 4): alle Dokumentationsfelder gefüllt
    row = 4
    assert ws.cell(row=row, column=headers["Geprüft am"]).value
    assert ws.cell(row=row, column=headers["Letzte Zeichnungsänderung"]
                   ).value == "2026-01-12"
    procs = ws.cell(row=row, column=headers["Fertigungsverfahren"]).value
    assert "Spanende Bearbeitung" in procs and "Gewindefertigung" in procs
    shot = ws.cell(row=row, column=headers["Screenshot"])
    assert shot.hyperlink is not None
    # Schweißkonsole (Zeile 2): Schweißen erkannt
    procs2 = ws.cell(row=2, column=headers["Fertigungsverfahren"]).value
    assert "Schweißen" in procs2
